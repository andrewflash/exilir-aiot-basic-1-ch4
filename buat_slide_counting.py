"""Extend v8 with three editable counting practice slides, without rebuilding it."""
from pathlib import Path
import ast,re,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parent
# Reuse only layout functions; importing the prior builder would rebuild old decks.
tree=ast.parse((R/'buat_slide_lab.py').read_text(encoding='utf-8'))
for node in tree.body:
    if isinstance(node,ast.FunctionDef) and node.name in ['box','add']:
        exec(compile(ast.Module(body=[node],type_ignores=[]),'<layout>','exec'))
slides=[
dict(title='Lab 9 • Hitung langkah langsung di ESP32',points=[
'Upload Exilir_Step_Counter; buka Serial Monitor 115200.',
'Kirim r → diam 3 detik → lakukan 20 langkah / stepping.',
'Bandingkan kolom steps dengan hitungan pasangan.',
'Gyroscope tidak dipakai; parameter ada di StepDetector.h.'],
note='Sketch dan header harus satu folder. SDA2/SCL3, MPU6050 0x68; lanjutkan setup Chapter 3. Jika memakai USB, stepping di tempat dan beri label sesuai aktivitas. Reset sebelum setiap percobaan. Hitungan ini estimasi kandidat langkah, bukan ground truth. Belum diuji pada kit fisik.',code=None),
dict(title='Lab 10 • Aturan threshold dalam kode',points=[
'Magnitude → moving average 5 sampel → kurangi baseline.',
'Terima puncak lokal > 1,2 m/s² dengan jeda ≥300 ms.',
'Sinyal harus turun <0,3 m/s² sebelum kandidat berikutnya.',
'Coba threshold 0,8 / 1,2 / 1,8; uji sesi baru.'],
note='Baseline median 3 detik diam. Puncak sebelumnya dikonfirmasi dengan sampel berikutnya. τ_min adalah interval minimum antara dua puncak diterima. Potongan di slide hanya pemakaian class; implementasi lengkap ada pada StepDetector.h. Ubah satu parameter per percobaan, catat false positive/negative. Jangan mengklaim akurasi dari sinyal sintetis.',
code='bool accepted = detector.update(t_ms, ax, ay, az);\n// di dalam detector: peak + threshold + hysteresis + jeda\nSerial.println(detector.count);\n// reset sesi: detector.reset(); mulai ulang t_ms dari 0'),
dict(title='Lab 11 • Counter + CSV dari handphone',points=[
'Upload Exilir_SoftAP_Logger; nyalakan dengan power bank.',
'Wi-Fi Exilir-IMU-XXXX → buka http://192.168.4.1.',
'Start → diam 3 detik → bergerak → Stop → unduh CSV + metadata.',
'Catat hitungan HP, manual, dan hasil Python; jelaskan selisih.'],
note='Password exilir12345. Counter diproses di ESP32 target50Hz; tampilan HP diperbarui sekitar1detik. Start mereset count; Stop mempertahankan sampai Start/reset berikutnya. Metadata berisi steps, threshold, baseline dan gap. CSV tetap kanal sensor untuk analisis ulang. Default23detik=3baseline+20aktivitas; samakan interval referensi. Data hilang saat mati/reset. Koneksi power bank dan Wi-Fi masih perlu pengujian fisik. Panduan: lab/Exilir_Step_Counter/README.md.',code=None)]
prs=Presentation(R.parent/'Anatomi_Langkah_Kaki_Exilir_v8.pptx')
insert=next(i for i,sl in enumerate(prs.slides) if any(sh.has_text_frame and 'Cek pemahaman sebelum pulang' in sh.text for sh in sl.shapes))
for i,d in enumerate(slides):
    add(prs,d,insert+i+1);ids=prs.slides._sldIdLst;sid=ids[-1];ids.remove(sid);ids.insert(insert+i,sid)
for i,sl in enumerate(prs.slides,1):
    sl.notes_slide.notes_text_frame.text=sl.notes_slide.notes_text_frame.text.replace('slide lampiran 53','slide lampiran 56')
    for sh in sl.shapes:
        if not sh.has_text_frame or sh.top<Inches(6.95):continue
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                if re.fullmatch(r'\d{2}',run.text.strip()):run.text=f'{i:02}'
                elif 'AIoT BASIC |' in run.text:run.text=re.sub(r'(AIoT BASIC \| )\d+',lambda m:m[1]+f'{i:02}',run.text)
out=R.parent/'Anatomi_Langkah_Kaki_Exilir_v9.pptx';prs.save(out)
single=Presentation();single.slide_width=Inches(13.333);single.slide_height=Inches(7.5)
for i,d in enumerate(slides,1):add(single,d,i)
single.save(R/'Hands_On_Counting_Arduino.pptx')
md=['# Tambahan hands-on counting',f'Deck v9: slide {insert+1}–{insert+3}.']
for d in slides:
    md.extend(['\n## '+d['title'],*['- '+p for p in d['points']],'\nCatatan: '+d['note']])
    if d['code']:md+=['\n```cpp\n'+d['code']+'\n```']
(R/'Konten_Slide_Counting.md').write_text('\n'.join(md),encoding='utf-8')
assert len(Presentation(out).slides)==56
print(f'PASS: v9 56 slides; added {insert+1}-{insert+3}, plus separate 3-slide deck.')
