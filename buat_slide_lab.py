from pathlib import Path
import re, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parent
slides=[]
def s(title,points,note,code=None):slides.append(dict(title=title,points=points,note=note,code=code))
s('Hands-on lab: dari IMU ke hitungan langkah',[
'Bekerja berpasangan: satu peserta bergerak, satu merekam dan menghitung.',
'Hasil akhir: CSV + metadata + grafik + tabel evaluasi.',
'Urutan: verifikasi sensor → rekam → plot → deteksi → evaluasi.',
'Alokasi 45–60 menit; siapkan Python dan Arduino sebelum kelas.'],
'Lab memakai firmware USB sederhana. Untuk setup bertambat lakukan stepping di tempat dan beri label sesuai aktivitas aktual. Jalan di lintasan bebas menjadi latihan setelah logger nirkabel tersedia. Tidak melakukan training ML dalam lab ini.')
s('Lab 1 • Lanjutkan kode pembacaan IMU Chapter 3',[
'Sketch baru: lab/Exilir_IMU_Logger/Exilir_IMU_Logger.ino',
'ESP32-C5, SDA GPIO2, SCL GPIO3, alamat 0x68 sesuai proyek.',
'Install Adafruit MPU6050, Unified Sensor, dan BusIO.',
'Pilih board/port aktual, compile lalu upload. WHO_AM_I harus 0x68.'],
'Referensi: Chapter_3/Baca_IMU_Library. Kode sebelumnya mencampur nama MPU6500/MPU6050. Sketch ini secara eksplisit untuk MPU6050; jika sensor berbeda, gunakan driver yang benar. Baud 115200. Matikan Serial Monitor saat logger Python memakai port. Compile diuji terpisah dari upload; belum uji hardware.')
s('Lab 2 • Konfigurasi akuisisi yang konsisten',[
'Target 50 Hz; pembacaan tiap 20 ms.',
'Rentang ±8 g; filter hardware 10 Hz.',
'Rekam akselerasi dan gyro; detector awal memakai akselerasi.',
'Saat diam: magnitude sekitar 9,81 m/s²; gyro sekitar nol.'],
'Adafruit menghasilkan akselerasi dalam m/s² dan kecepatan rotasi dalam rad/s. Timestamp dibuat di perangkat pada awal pembacaan. Polling tetap memiliki jitter, sehingga interval aktual harus diperiksa. Snippet adalah konfigurasi; gunakan sketch lengkap untuk setup dan loop.',
'''Wire.begin(2, 3);
mpu.begin(0x68, &Wire);
mpu.setAccelerometerRange(MPU6050_RANGE_8_G);
mpu.setGyroRange(MPU6050_RANGE_500_DEG);
mpu.setFilterBandwidth(MPU6050_BAND_10_HZ);''')
s('Lab 3 • Apa yang dilakukan loop logger?',[
'Tunggu perintah s untuk memulai sesi.',
'Baca hanya ketika jadwal 20 ms tiba.',
'Catat timestamp aktual dan keluarkan satu baris CSV.',
'Akhiri setelah 23 detik; jangan membuat sampel palsu saat terlambat.'],
'Ini potongan inti, bukan sketch mandiri; penjadwalan dan penanda BEGIN/ACTIVITY/END ada di file .ino lengkap. Header: timestamp,accX,accY,accZ,gyroX,gyroY,gyroZ. Serial.print menulis kanal lainnya dalam sketch lengkap. Timestamp dari micros() dikonversi menjadi milidetik relatif.',
'''sensors_event_t a, g, temperature;
mpu.getEvent(&a, &g, &temperature);
float t_ms = (micros() - started) / 1000.0;
// satu baris: waktu, ax, ay, az, gx, gy, gz
// a.acceleration.x  -> m/s2
// g.gyro.x          -> rad/s''')
s('Lab 4 • Simpan rekaman ke CSV di laptop',[
'Jalankan perintah dari folder Slide/lab.',
'Tutup Serial Monitor; ganti COM7 dengan port kit.',
'DIAM: baseline 3 detik. MULAI: aktivitas dan hitungan 20 detik.',
'Output disimpan ke folder rekaman; gunakan ID sesi baru.'],
'Pasang dependensi sekali sebelum kelas. Script mengirim s secara otomatis. Cue terminal memiliki latensi USB dan reaksi pengamat. Untuk referensi waktu yang presisi gunakan video tersinkron. Contoh ini memakai stepping_in_place, karena firmware USB masih bertambat.',
'''python -m pip install -r requirements.txt
python -m serial.tools.list_ports
python rekam_serial.py --port COM7 --subject P01 --session S01 --label stepping_in_place''')
s('Lab 5 • Periksa tiga berkas hasil',[
'*.csv: 23 detik penuh, enam kanal + timestamp.',
'*.activity.csv: 20 detik aktivitas, hanya tiga kanal akselerasi.',
'*.json: subjek, sesi, label, posisi, dan kualitas sampling.',
'Isi reference_steps sesuai pengamat; baseline tidak ikut dihitung.'],
'CSV aktivitas membuang bagian baseline dan mengembalikan timestamp awal ke nol; cocok untuk latihan import berikutnya setelah timing diperiksa. Label tidak menjadi kolom numerik sensor. Jangan menganalisis .activity.csv dengan analyzer baseline. Frekuensi nyata dihitung dari selisih timestamp; logger menandai interval di luar 18–22 ms. Satuan dan kanal harus tetap konsisten.')
s('Lab 6 • Plot sinyal dan jalankan peak detector',[
'Gunakan CSV mentah yang masih memiliki baseline.',
'Median baseline → magnitude → moving average lima sampel.',
'Awal: T_high=1,2; T_low=0,3 m/s²; τ_min=0,30 detik.',
'Buka *.plot.png dan *.hasil.json; tandai puncak salah atau hilang.'],
'Ganti 40 dengan hitungan aktual pengamat. Analisis menghasilkan count dan error relatif jika reference >0. Untuk reference 0 gunakan false steps/menit. Analyzer menolak timestamp terlalu tidak teratur; periksa logger, jangan memaksa lanjut. Kelima sampel filtering pada 50Hz setara jendela100ms; ada delay yang perlu dipahami.',
'''python analisis_langkah.py rekaman/stepping_in_place_P01_S01.csv --reference 40

# Setelah tuning pada data latihan:
# tambahkan --high 1.8 --min-interval 0.4''')
s('Lab 7 • Uji lima kondisi, ulangi tiga kali',[
'Diam; stepping pelan; stepping biasa; stepping cepat.',
'Goyang tanpa berjalan sebagai contoh aktivitas pengganggu.',
'Bekukan parameter setelah tuning; uji sesi baru dan orang kedua.',
'Catat posisi sensor dan aktivitas aktual pada setiap rekaman.'],
'Jika melakukan jalan di lintasan dengan setup yang memadai, ganti label stepping menjadi walking; jangan menyamakan kedua domain. Jangan membandingkan parameter hanya pada satu rekaman. Gunakan template_hasil.csv untuk jumlah manual/prediksi dan konfigurasi. Untuk goyangan setelah sensor dilepas, posisi menjadi hand; metadata harus mencerminkannya.')
s('Lab 8 • Apa yang harus dikumpulkan?',[
'CSV, metadata, dan satu grafik yang sudah diberi penjelasan.',
'Tabel N_ref, N_pred, error, threshold, jeda, dan filter.',
'Contoh kasus berhasil serta kegagalan yang ditemukan.',
'Kesimpulan: aturan cukup, perlu perbaikan, atau layak mencoba ML?'],
'Kriteria: data lengkap, timestamp masuk akal, baseline benar-benar diam, hitungan pada interval sama, set uji tidak dipakai untuk tuning. Tidak menetapkan persentase akurasi universal. Peserta harus menjelaskan alasan salah hitung; ML adalah hipotesis perbaikan yang perlu diuji, bukan jawaban otomatis.')
s('Bonus • Ubah logger menjadi nirkabel',[
'Referensi koneksi: Chapter_2/MQTT_Example.',
'Sampling → ring buffer → batch → MQTT → CSV laptop.',
'Tambahkan session_id, sequence, dan timestamp sensor.',
'Setelah logger bekerja tanpa USB, gunakan power bank.'],
'Bonus masih rancangan pengembangan, bukan firmware Wi-Fi siap pakai. Sketch lab USB tidak menyimpan data ketika hanya diberi power bank. Pisahkan akuisisi dari jaringan, deteksi buffer overflow dan paket hilang/duplikat. Uji end-to-end sebelum berjalan bebas. Di sesi selanjutnya file aktivitas yang sudah diperiksa dapat diimpor ke Edge Impulse.')

def box(sl,x,y,w,h,value,size=22,color='162C40',bold=False,font='Aptos'):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for j,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(12)
    return sh
def add(prs,d,num):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string('F6F8FA')
    box(sl,.6,.35,12,.3,'HANDS-ON LAB • CHAPTER 4',11,'147D92',True)
    box(sl,.6,.9,12.1,.95,d['title'],29,bold=True)
    if d['code']:
        box(sl,.65,2,11.9,2.45,'\n'.join('• '+p for p in d['points']),20)
        box(sl,.8,4.6,11.9,2.3,d['code'],15,'147D92',font='Consolas')
    else:box(sl,.7,2.15,11.8,4.4,'\n'.join('• '+p for p in d['points']),25)
    box(sl,.6,7.07,12,.25,f'EXILIR • AIoT BASIC | {num:02} • Panduan lengkap: lab/README.md',10,'657786')
    sl.notes_slide.notes_text_frame.text=d['note']
    return sl
prs=Presentation(R.parent/'Anatomi_Langkah_Kaki_Exilir_v7.pptx')
insert=next(i for i,sl in enumerate(prs.slides) if any(sh.has_text_frame and 'Cek pemahaman sebelum pulang' in sh.text for sh in sl.shapes))
old=len(prs.slides)
for i,d in enumerate(slides):
    add(prs,d,insert+i+1);ids=prs.slides._sldIdLst;sid=ids[-1];ids.remove(sid);ids.insert(insert+i,sid)
for i,sl in enumerate(prs.slides,1):
    if i<=insert+len(slides):continue
    for sh in sl.shapes:
        if not sh.has_text_frame or sh.top<Inches(6.95):continue
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                if re.fullmatch(r'\d{2}',run.text.strip()):run.text=f'{i:02}'
                elif 'AIoT BASIC |' in run.text:run.text=re.sub(r'(AIoT BASIC \| )\d+',lambda m:m[1]+f'{i:02}',run.text)
# Keep references to the moved WISDM appendix correct.
for sl in prs.slides:
    tf=sl.notes_slide.notes_text_frame;tf.text=tf.text.replace('slide lampiran 43','slide lampiran 53')
out=R.parent/'Anatomi_Langkah_Kaki_Exilir_v8.pptx';prs.save(out)
single=Presentation();single.slide_width=Inches(13.333);single.slide_height=Inches(7.5)
for i,d in enumerate(slides,1):add(single,d,i)
single.save(R/'Hands_On_Lab_Exilir.pptx')
md=['# Konten slide hands-on lab',f'Pada deck v8: slide {insert+1}–{insert+len(slides)}; total {old+len(slides)} slide.']
for i,d in enumerate(slides,1):
    md.extend([f'\n## {i}. {d["title"]}',*['- '+p for p in d['points']]])
    if d['code']:md+=['\n```text\n'+d['code']+'\n```']
    md+=['\nCatatan pengajar: '+d['note']]
(R/'Konten_Slide_Lab.md').write_text('\n'.join(md),encoding='utf-8')
assert len(Presentation(out).slides)==53
(R/'slides_lab.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Generated 10 lab slides; full deck 53 slides. New slides {insert+1}-{insert+len(slides)}.')
