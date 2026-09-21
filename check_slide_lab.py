from pathlib import Path
from pptx import Presentation
from PIL import ImageFont
p=Presentation(Path(__file__).parent/'Hands_On_Lab_Exilir.pptx');issues=[]
for i,s in enumerate(p.slides,1):
    for sh in s.shapes:
        if not sh.has_text_frame:continue
        height=0
        for para in sh.text_frame.paragraphs:
            pt=para.font.size.pt if para.font.size else 22
            font='consola.ttf' if para.font.name=='Consolas' else 'arial.ttf'
            f=ImageFont.truetype('C:/Windows/Fonts/'+font,round(pt*96/72))
            width=sh.width/914400*96;lines=1;line=''
            for word in para.text.split():
                candidate=(line+' '+word).strip()
                if f.getlength(candidate)>width:lines+=1;line=word
                else:line=candidate
            height+=lines*pt*1.18+12
        height-=12
        if height>sh.height/914400*72+5:issues.append((i,sh.text[:35],round(height),round(sh.height/914400*72)))
print('Estimated overflow:',issues)
assert not issues
